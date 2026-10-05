from io import BytesIO
from tempfile import TemporaryDirectory
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.files.base import ContentFile
from django.test import override_settings
from rest_framework.test import APITestCase
from rest_framework_simplejwt.tokens import RefreshToken
from docx import Document

from .models import Application, Company, Job, PlatformSettings, SavedJob


User = get_user_model()


class BackendApiTests(APITestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            email="candidate@example.com",
            password="Test-password-471!",
            name="Candidate",
        )
        self.admin = User.objects.create_superuser(
            email="admin@example.com",
            password="Test-password-472!",
            name="Administrator",
        )
        self.job = Job.objects.create(
            title="Frontend Engineer",
            company="Example Co",
            location="Remote",
            description="React and JavaScript",
        )

    def authenticate_as(self, user):
        token = str(RefreshToken.for_user(user).access_token)
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {token}")

    def test_public_jobs_and_companies(self):
        Company.objects.create(
            name="Example Co",
            industry="Technology",
            location="Remote",
        )

        jobs_response = self.client.get("/jobs/")
        companies_response = self.client.get("/companies/")

        self.assertEqual(jobs_response.status_code, 200)
        self.assertEqual(jobs_response.data[0]["title"], "Frontend Engineer")
        self.assertEqual(companies_response.status_code, 200)
        self.assertEqual(companies_response.data[0]["jobs"], 1)

    def test_registration_cannot_assign_admin_role(self):
        response = self.client.post(
            "/auth/register",
            {
                "name": "New Candidate",
                "email": "new@example.com",
                "password": "Valid-password-993!",
                "role": "admin",
            },
            format="json",
        )

        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.data["user"]["role"], User.Role.USER)
        self.assertNotIn("password", response.data["user"])

    def test_login_returns_the_user_record_needed_by_the_frontend(self):
        response = self.client.post(
            "/auth/login",
            {
                "email": self.user.email,
                "password": "Test-password-471!",
            },
            format="json",
        )

        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.data["access_token"])
        self.assertEqual(response.data["user"]["id"], self.user.id)
        self.assertEqual(response.data["user"]["name"], self.user.name)
        self.assertEqual(response.data["user"]["role"], User.Role.USER)

    def test_admin_only_job_management_and_application_lifecycle(self):
        self.authenticate_as(self.user)
        denied = self.client.post(
            "/jobs/",
            {"title": "Restricted", "company": "Example", "location": "Remote"},
            format="json",
        )
        self.assertEqual(denied.status_code, 403)

        application_response = self.client.post(
            "/applications/",
            {"job_id": self.job.id, "cover_letter": "I am interested."},
            format="json",
        )
        self.assertEqual(application_response.status_code, 201)
        self.assertEqual(application_response.data["job"], self.job.title)
        self.assertEqual(application_response.data["job_id"], self.job.id)
        self.assertEqual(application_response.data["applicant"], self.user.name)

        self.authenticate_as(self.admin)
        update_response = self.client.put(
            f"/applications/admin/{application_response.data['id']}",
            {"status": Application.Status.SHORTLISTED},
            format="json",
        )
        self.assertEqual(update_response.status_code, 200)
        self.assertEqual(update_response.data["status"], Application.Status.SHORTLISTED)

        delete_response = self.client.delete(
            f"/applications/admin/{application_response.data['id']}"
        )
        self.assertEqual(delete_response.status_code, 204)

    def test_user_saved_jobs_are_scoped_to_the_authenticated_user(self):
        self.authenticate_as(self.user)
        create_response = self.client.post(
            "/saved-jobs/",
            {"job_id": self.job.id},
            format="json",
        )
        self.assertEqual(create_response.status_code, 201)
        self.assertEqual(create_response.data["job_id"], self.job.id)
        self.assertEqual(create_response.data["saved_at"], create_response.data["created_at"])

        saved_job = SavedJob.objects.get(user=self.user, job=self.job)
        other_user = User.objects.create_user(
            email="other@example.com",
            password="Test-password-473!",
            name="Other Candidate",
        )
        self.authenticate_as(other_user)
        delete_response = self.client.delete(f"/saved-jobs/{saved_job.id}")
        self.assertEqual(delete_response.status_code, 404)
        self.assertTrue(SavedJob.objects.filter(pk=saved_job.id).exists())

    def test_admin_job_creation_and_settings_match_frontend_contract(self):
        self.authenticate_as(self.admin)
        PlatformSettings.objects.create(
            pk=1,
            site_name="JobAI",
            support_email="support@jobai.com",
            default_job_status=Job.Status.DRAFT,
        )
        job_response = self.client.post(
            "/jobs",
            {
                "title": "Backend Engineer",
                "company": "Example Co",
                "location": "Remote",
            },
            format="json",
        )
        self.assertEqual(job_response.status_code, 201)
        self.assertEqual(job_response.data["status"], Job.Status.DRAFT)

        settings_response = self.client.get("/settings/")
        self.assertEqual(settings_response.status_code, 200)
        self.assertEqual(settings_response.data["profile"]["name"], self.admin.name)
        self.assertEqual(settings_response.data["platform"]["platformName"], "JobAI")

        profile_response = self.client.put(
            "/settings/profile",
            {"name": "Updated Admin", "email": self.admin.email, "phone": "555-0100"},
            format="json",
        )
        self.assertEqual(profile_response.status_code, 200)
        self.assertEqual(profile_response.data["phone"], "555-0100")

    @patch("api.ai._generate_json")
    def test_resume_analysis_returns_gemini_feedback_for_uploaded_docx(self, generate_json):
        generate_json.return_value = {
            "ats_score": 86,
            "summary": "A focused engineering resume with relevant skills.",
            "strengths": ["Clear project experience"],
            "improvements": ["Add measurable outcomes"],
            "detected_skills": ["Python", "Django"],
            "detected_sections": ["skills", "projects"],
        }
        document = Document()
        document.add_paragraph("Python Django project experience and skills")
        buffer = BytesIO()
        document.save(buffer)

        with TemporaryDirectory() as media_root, override_settings(MEDIA_ROOT=media_root):
            self.user.resume.save(
                "candidate.docx",
                ContentFile(buffer.getvalue()),
                save=True,
            )
            self.authenticate_as(self.user)
            response = self.client.get("/ai/resume/analyze")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["filename"], "candidate.docx")
        self.assertEqual(response.data["analysis"]["ats_score"], 86)
        self.assertEqual(response.data["analysis"]["word_count"], 6)
        self.assertEqual(response.data["analysis"]["detected_skills"], ["Python", "Django"])
        generate_json.assert_called_once()

    @patch("api.ai._generate_json")
    def test_ai_job_matches_rank_only_known_jobs(self, generate_json):
        second_job = Job.objects.create(
            title="Data Engineer",
            company="Data Co",
            location="Remote",
            description="Python, SQL, and data pipeline development",
        )
        generate_json.return_value = {
            "recommendations": [
                {
                    "job_id": second_job.id,
                    "match_score": 91,
                    "matched_skills": ["Python", "SQL"],
                    "missing_skills": ["Airflow"],
                    "fit_summary": "Strong match based on Python and SQL projects.",
                },
                {
                    "job_id": 999999,
                    "match_score": 100,
                    "matched_skills": ["Unknown"],
                    "missing_skills": [],
                    "fit_summary": "Unknown job.",
                },
            ]
        }
        document = Document()
        document.add_paragraph("Python SQL project experience")
        buffer = BytesIO()
        document.save(buffer)

        with TemporaryDirectory() as media_root, override_settings(MEDIA_ROOT=media_root):
            self.user.resume.save(
                "candidate.docx",
                ContentFile(buffer.getvalue()),
                save=True,
            )
            self.authenticate_as(self.user)
            response = self.client.get("/ai/job-matches/")
            recommendations_response = self.client.get("/ai/recommendations/")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.data["matches"]), 1)
        self.assertEqual(response.data["matches"][0]["job_id"], second_job.id)
        self.assertEqual(response.data["matches"][0]["match"]["match_score"], 91)
        self.assertEqual(
            recommendations_response.data["recommendations"][0]["job_id"],
            second_job.id,
        )
        self.assertNotIn(
            "description",
            recommendations_response.data["recommendations"][0],
        )
        self.assertEqual(generate_json.call_count, 2)

    @patch("api.ai._generate_json")
    def test_ai_job_match_analyzes_the_selected_active_job(self, generate_json):
        generate_json.return_value = {
            "match_score": 88,
            "matched_skills": ["React", "JavaScript"],
            "missing_skills": ["TypeScript"],
            "fit_summary": "Strong frontend fit based on the resume projects.",
        }
        document = Document()
        document.add_paragraph("React JavaScript frontend project experience")
        buffer = BytesIO()
        document.save(buffer)

        with TemporaryDirectory() as media_root, override_settings(MEDIA_ROOT=media_root):
            self.user.resume.save(
                "candidate.docx",
                ContentFile(buffer.getvalue()),
                save=True,
            )
            self.authenticate_as(self.user)
            response = self.client.get(f"/ai/job-match/{self.job.id}")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["match"]["match_score"], 88)
        self.assertEqual(
            response.data["match"]["matched_skills"],
            ["React", "JavaScript"],
        )
        generate_json.assert_called_once()

    @override_settings(GEMINI_API_KEY="")
    def test_ai_endpoints_require_a_resume_before_analysis(self):
        self.authenticate_as(self.user)

        response = self.client.get("/ai/recommendations/")

        self.assertEqual(response.status_code, 400)
        self.assertIn("Upload a PDF or DOCX resume", response.data["detail"])

    def test_invalid_resume_document_returns_a_validation_error(self):
        with TemporaryDirectory() as media_root, override_settings(MEDIA_ROOT=media_root):
            self.user.resume.save(
                "invalid.docx",
                ContentFile(b"not a valid document"),
                save=True,
            )
            self.authenticate_as(self.user)
            response = self.client.get("/ai/resume/analyze")

        self.assertEqual(response.status_code, 400)
        self.assertIn("could not be read", response.data["detail"])

    @override_settings(GEMINI_API_KEY="")
    def test_ai_endpoint_explains_when_gemini_is_not_configured(self):
        document = Document()
        document.add_paragraph("Python Django resume")
        buffer = BytesIO()
        document.save(buffer)

        with TemporaryDirectory() as media_root, override_settings(MEDIA_ROOT=media_root):
            self.user.resume.save(
                "candidate.docx",
                ContentFile(buffer.getvalue()),
                save=True,
            )
            self.authenticate_as(self.user)
            response = self.client.get("/ai/recommendations/")

        self.assertEqual(response.status_code, 503)
        self.assertIn("GEMINI_API_KEY", response.data["detail"])

    @override_settings(
        GEMINI_API_KEY="test-api-key",
        GEMINI_MODEL="gemini-test-model",
    )
    @patch("api.ai.requests.post")
    def test_gemini_client_sends_key_and_parses_json_response(self, post):
        post.return_value.ok = True
        post.return_value.status_code = 200
        post.return_value.json.return_value = {
            "candidates": [
                {"content": {"parts": [{"text": '{"ready": true}'}]}}
            ]
        }

        from .ai import _generate_json

        result = _generate_json("Return JSON.")

        self.assertEqual(result, {"ready": True})
        self.assertEqual(post.call_args.kwargs["headers"]["x-goog-api-key"], "test-api-key")
        self.assertIn("gemini-test-model", post.call_args.args[0])
